#!/usr/bin/env python3
"""Genera las plantillas .docx base (relleno) en /plantillas.

Uso:  pip install python-docx && python3 infra/generar_plantillas_base.py

Cada plantilla trae membrete (encabezado) y pie con variables de la firma,
y un cuerpo de EJEMPLO con variables {{grupo.campo}} que el sistema
reemplaza con los datos del proceso. El contenido real lo define la abogada:
basta con abrir el .docx en Word, editarlo y volver a subirlo desde la vista
Plantillas (conservando las variables {{...}}).
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DEST = os.path.join(os.path.dirname(__file__), '..', 'plantillas')
AZUL = RGBColor(0x2C, 0x3E, 0x50)

def base():
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2)
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

    h = s.header
    p = h.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('{{empresa.nombre}}'); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = AZUL
    p2 = h.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run('NIT {{empresa.nit}}  ·  {{empresa.ciudad}}'); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    f = s.footer
    fp = f.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('{{empresa.direccion}}  ·  Tel. {{empresa.telefono}}  ·  {{empresa.email}}')
    r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    return d

def p(d, txt='', bold=False, align=None, size=None, italic=False, space=6):
    para = d.add_paragraph()
    if align == 'c': para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == 'j': para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if align == 'r': para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.paragraph_format.space_after = Pt(space)
    r = para.add_run(txt); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return para

def borrador(d):
    p(d, '[BORRADOR — contenido de ejemplo pendiente de definir con la abogada]', italic=True, size=9, space=12)

def firma(d):
    p(d, '', space=24)
    p(d, '______________________________')
    p(d, '{{abogado.nombre}}', bold=True, space=0)
    p(d, '{{abogado.cargo}}', space=0)
    p(d, '{{abogado.email}}  ·  {{abogado.telefono}}', size=9)

def salvar(d, nombre):
    os.makedirs(DEST, exist_ok=True)
    d.save(os.path.join(DEST, nombre))
    print('OK', nombre)

# 1. Poder especial
d = base(); borrador(d)
p(d, '{{fecha.lugar_y_fecha}}', align='r', space=14)
p(d, 'Señor(a)\n{{proceso.despacho}}\n{{proceso.municipio}}\nE. S. D.', space=12)
p(d, 'Referencia: Proceso {{proceso.tipo}} — Radicado {{proceso.radicado}}', bold=True, space=12)
p(d, 'Asunto: PODER ESPECIAL', bold=True, space=12)
p(d, '{{cliente.nombre_completo}}, mayor de edad, identificado(a) con {{cliente.tipo_identificacion}} No. {{cliente.numero_identificacion}}, '
     'domiciliado(a) en {{cliente.direccion}}, manifiesto que confiero PODER ESPECIAL, amplio y suficiente, al abogado(a) '
     '{{abogado.nombre}}, para que en mi nombre y representación adelante y lleve hasta su terminación el proceso de la referencia.', align='j')
p(d, 'El apoderado queda facultado para recibir, desistir, transigir, conciliar, sustituir, reasumir, y en general para todas las '
     'actuaciones necesarias para el cabal cumplimiento de este mandato.', align='j', space=12)
p(d, 'Atentamente,', space=24)
p(d, '______________________________')
p(d, '{{cliente.nombre_completo}}', bold=True, space=0)
p(d, '{{cliente.tipo_identificacion}} No. {{cliente.numero_identificacion}}', space=14)
p(d, 'Acepto,', space=18)
p(d, '______________________________')
p(d, '{{abogado.nombre}}', bold=True, space=0)
salvar(d, 'poder_especial.docx')

# 2. Demanda
d = base(); borrador(d)
p(d, '{{fecha.lugar_y_fecha}}', align='r', space=14)
p(d, 'Señor(a) Juez\n{{proceso.despacho}}\n{{proceso.municipio}}\nE. S. D.', space=12)
p(d, 'Referencia: DEMANDA — {{proceso.tipo}}', bold=True, space=0)
p(d, 'Demandante: {{cliente.nombre_completo}} ({{cliente.tipo_identificacion}} {{cliente.numero_identificacion}})', space=0)
p(d, 'Demandado: {{proceso.entidad}}', space=12)
p(d, '{{abogado.nombre}}, abogado(a) en ejercicio, obrando como apoderado(a) de {{cliente.nombre_completo}}, presento demanda en los siguientes términos:', align='j', space=12)
p(d, 'I. HECHOS', bold=True)
p(d, '1. {{proceso.descripcion}}', align='j')
p(d, '2. [Hecho 2]', align='j'); p(d, '3. [Hecho 3]', align='j', space=12)
p(d, 'II. PRETENSIONES', bold=True)
p(d, '1. [Pretensión 1]', align='j'); p(d, '2. [Pretensión 2]', align='j', space=12)
p(d, 'III. FUNDAMENTOS DE DERECHO', bold=True)
p(d, '[Normas y jurisprudencia aplicables]', align='j', space=12)
p(d, 'IV. PRUEBAS Y ANEXOS', bold=True)
p(d, '[Relación de pruebas y anexos]', align='j', space=12)
p(d, 'V. NOTIFICACIONES', bold=True)
p(d, 'El demandante recibe notificaciones en {{cliente.direccion}} y {{cliente.email}}. El suscrito, en {{abogado.email}}.', align='j')
firma(d); salvar(d, 'demanda.docx')

# 3. Contestación de demanda
d = base(); borrador(d)
p(d, '{{fecha.lugar_y_fecha}}', align='r', space=14)
p(d, 'Señor(a) Juez\n{{proceso.despacho}}\n{{proceso.municipio}}\nE. S. D.', space=12)
p(d, 'Referencia: CONTESTACIÓN DE LA DEMANDA', bold=True, space=0)
p(d, 'Proceso: {{proceso.tipo}} — Radicado {{proceso.radicado}}', space=12)
p(d, '{{abogado.nombre}}, obrando como apoderado(a) de {{cliente.nombre_completo}}, me permito contestar la demanda de la referencia:', align='j', space=12)
p(d, 'I. PRONUNCIAMIENTO SOBRE LOS HECHOS', bold=True); p(d, '[Al hecho 1: ...]', align='j', space=12)
p(d, 'II. PRONUNCIAMIENTO SOBRE LAS PRETENSIONES', bold=True); p(d, '[Me opongo / Acepto ...]', align='j', space=12)
p(d, 'III. EXCEPCIONES', bold=True); p(d, '[Excepciones de mérito y previas]', align='j', space=12)
p(d, 'IV. PRUEBAS', bold=True); p(d, '[Relación de pruebas]', align='j')
firma(d); salvar(d, 'contestacion_demanda.docx')

# 4. Notificación
d = base(); borrador(d)
p(d, '{{fecha.lugar_y_fecha}}', align='r', space=14)
p(d, 'Señores\n{{proceso.entidad}}\nCiudad', space=12)
p(d, 'Asunto: NOTIFICACIÓN — Proceso {{proceso.tipo}}, Radicado {{proceso.radicado}}', bold=True, space=12)
p(d, 'Por medio de la presente se notifica a {{proceso.entidad}} de la existencia del proceso de la referencia, que cursa en {{proceso.despacho}}, '
     'promovido por {{cliente.nombre_completo}} ({{cliente.tipo_identificacion}} {{cliente.numero_identificacion}}).', align='j')
p(d, 'Última actuación registrada: {{proceso.ultima_actuacion}} ({{proceso.ultima_actuacion_fecha}}).', align='j', space=12)
p(d, '[Texto de la notificación pendiente de definir]', align='j')
firma(d); salvar(d, 'notificacion.docx')

# 5. Derecho de petición
d = base(); borrador(d)
p(d, '{{fecha.lugar_y_fecha}}', align='r', space=14)
p(d, 'Señores\n{{proceso.entidad}}\nCiudad', space=12)
p(d, 'Asunto: DERECHO DE PETICIÓN', bold=True, space=12)
p(d, '{{cliente.nombre_completo}}, identificado(a) con {{cliente.tipo_identificacion}} No. {{cliente.numero_identificacion}}, actuando por intermedio de mi apoderado(a) '
     '{{abogado.nombre}}, en ejercicio del derecho de petición (art. 23 de la Constitución Política), solicito:', align='j', space=12)
p(d, '1. [Petición 1]', align='j'); p(d, '2. [Petición 2]', align='j', space=12)
p(d, 'Hechos que fundamentan la petición:', bold=True); p(d, '{{proceso.descripcion}}', align='j', space=12)
p(d, 'Notificaciones: {{cliente.direccion}} · {{cliente.email}} · {{abogado.email}}.', align='j')
firma(d); salvar(d, 'derecho_peticion.docx')

# 6. Memorial
d = base(); borrador(d)
p(d, '{{fecha.lugar_y_fecha}}', align='r', space=14)
p(d, 'Señor(a) Juez\n{{proceso.despacho}}\n{{proceso.municipio}}\nE. S. D.', space=12)
p(d, 'Referencia: Proceso {{proceso.tipo}} — Radicado {{proceso.radicado}}', bold=True, space=0)
p(d, 'Demandante / Interesado: {{cliente.nombre_completo}}', space=12)
p(d, 'Asunto: [Asunto del memorial]', bold=True, space=12)
p(d, '{{abogado.nombre}}, apoderado(a) de {{cliente.nombre_completo}}, me permito manifestar al despacho lo siguiente:', align='j', space=12)
p(d, '[Contenido del memorial]', align='j')
firma(d); salvar(d, 'memorial.docx')
